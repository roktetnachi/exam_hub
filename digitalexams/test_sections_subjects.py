"""Sections (formerly "classes"), subjects that mix pupils from several sections, and their bulk imports."""
import io
import json

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase

from .models import School, SchoolClass, Subject, User


def csv_file(text, name='import.csv'):
    return SimpleUploadedFile(name, text.encode('utf-8'), content_type='text/csv')


class SectionsAndSubjectsTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name='SSC', slug='ssc')
        self.other = School.objects.create(name='Other', slug='other')
        self.admin = User.objects.create_user('ad', 'ad@x.edu', 'x', school=self.school, role='teacher', is_school_admin=True)
        self.teacher = User.objects.create_user('eco', 'eco@x.edu', 'x', school=self.school, role='teacher', first_name='Eco')
        self.sec_a = SchoolClass.objects.create(school=self.school, name='Pre-U 1 - A')
        self.sec_b = SchoolClass.objects.create(school=self.school, name='Pre-U 1 - B')
        self.a1 = User.objects.create_user('a1', 'a1@x.edu', 'x', school=self.school, first_name='Ana')
        self.a2 = User.objects.create_user('a2', 'a2@x.edu', 'x', school=self.school, first_name='Abe')
        self.b1 = User.objects.create_user('b1', 'b1@x.edu', 'x', school=self.school, first_name='Ben')
        self.loose = User.objects.create_user('n1', 'n1@x.edu', 'x', school=self.school, first_name='Nia')
        self.far = User.objects.create_user('f1', 'f1@x.edu', 'x', school=self.other, first_name='Far')
        self.sec_a.pupils.add(self.a1, self.a2)
        self.sec_b.pupils.add(self.b1)

    def post(self, url, data):
        return self.client.post('/s/ssc' + url, json.dumps(data), content_type='application/json')

    def upload(self, kind, text, commit=False, skip=False):
        data = {'file': csv_file(text), 'commit': '1' if commit else '0'}
        if skip:
            data['skipErrors'] = '1'
        return self.client.post(f'/s/ssc/api/bulk/{kind}/import/', data)

    # ---------------- pages & labels
    def test_pages_say_sections_and_link_subjects(self):
        self.client.force_login(self.admin)
        page = self.client.get('/s/ssc/classes/').content.decode()
        self.assertIn('<h1>Year groups</h1>', page)
        self.assertIn('Bulk add pupils to year groups', page)
        self.assertIn('/s/ssc/subjects/', page)
        self.assertEqual(self.client.get('/s/ssc/sections/').status_code, 200)
        self.assertEqual(self.client.get('/s/ssc/year-groups/').status_code, 200)
        subjects = self.client.get('/s/ssc/subjects/').content.decode()
        self.assertIn('Bulk add pupils to subjects', subjects)
        self.assertIn('Bulk import subjects', subjects)
        home = self.client.get('/s/ssc/').content.decode()
        self.assertIn('>Year groups</a>', home)
        self.assertIn('>Subjects</a>', home)

    # ---------------- subject API
    def test_admin_builds_subject_from_two_sections_and_teacher_can_assign(self):
        self.client.force_login(self.admin)
        r = self.post('/api/subjects/create/', {'name': 'Economics'})
        self.assertEqual(r.status_code, 201)
        sid = r.json()['id']
        self.assertEqual(self.post('/api/subjects/create/', {'name': 'economics'}).status_code, 400)
        r = self.post(f'/api/subjects/{sid}/update/', {'addPupilIds': [self.a1.id, self.b1.id, self.far.id],
                                                       'teacherIds': [self.teacher.id]})
        self.assertEqual(r.status_code, 200)
        body = r.json()
        self.assertEqual({p['email'] for p in body['pupils']}, {'a1@x.edu', 'b1@x.edu'})  # other school ignored
        self.assertEqual({tuple(p['sections']) for p in body['pupils']}, {('Pre-U 1 - A',), ('Pre-U 1 - B',)})
        self.post(f'/api/subjects/{sid}/update/', {'removePupilIds': [self.b1.id]})
        self.assertEqual(list(Subject.objects.get(id=sid).pupils.all()), [self.a1])

        # The subject teacher (no section of their own) can see and assign to its pupils only.
        self.client.force_login(self.teacher)
        emails = {p['email'] for p in self.client.get('/s/ssc/api/pupils/').json()['pupils']}
        self.assertEqual(emails, {'a1@x.edu'})
        mine = self.client.get('/s/ssc/api/subjects/').json()['subjects']
        self.assertEqual([s['name'] for s in mine], ['Economics'])
        self.assertEqual(self.post('/api/subjects/create/', {'name': 'Nope'}).status_code, 403)

    def test_subject_scoped_to_school(self):
        foreign = Subject.objects.create(school=self.other, name='Theirs')
        self.client.force_login(self.admin)
        self.assertEqual(self.post(f'/api/subjects/{foreign.id}/update/', {'name': 'Mine'}).status_code, 403)
        self.assertEqual(self.post(f'/api/subjects/{foreign.id}/delete/', {}).status_code, 403)
        self.assertEqual(self.client.get('/s/ssc/api/subjects/').json()['subjects'], [])

    # ---------------- bulk: subjects
    def test_bulk_import_subjects(self):
        self.client.force_login(self.admin)
        Subject.objects.create(school=self.school, name='Physics')
        text = 'Subject name,Teacher emails,Rename to\nEconomics,eco@x.edu,\nPhysics,,Physics H2\nMaths,nobody@x.edu,\n'
        r = self.upload('subjects', text)
        self.assertEqual(r.status_code, 200, r.content)
        self.assertEqual(r.json()['summary'], {'create': 1, 'update': 1, 'unchanged': 0, 'error': 1})
        r = self.upload('subjects', text, commit=True, skip=True)
        self.assertTrue(r.json()['committed'])
        self.assertEqual(sorted(Subject.objects.filter(school=self.school).values_list('name', flat=True)), ['Economics', 'Physics H2'])
        self.assertEqual(list(Subject.objects.get(name='Economics').teachers.all()), [self.teacher])

    # ---------------- bulk: pupils into sections
    def test_bulk_add_pupils_to_sections(self):
        self.client.force_login(self.admin)
        text = ('Email,Pupil name,Current section(s),Add to section\n'
                'n1@x.edu,Nia,,Pre-U 1 - B\n'
                'a1@x.edu,Ana,Pre-U 1 - A,Pre-U 1 - A\n'   # already there
                'b1@x.edu,Ben,Pre-U 1 - B,\n'               # blank → skipped
                'ghost@x.edu,Ghost,,Pre-U 1 - A\n'          # no account
                'n1@x.edu,Nia,,Pre-U 9\n')                  # unknown section
        r = self.upload('section_pupils', text)
        self.assertEqual(r.json()['summary'], {'create': 0, 'update': 1, 'unchanged': 1, 'error': 2})
        self.assertEqual(self.upload('section_pupils', text, commit=True).status_code, 400)  # errors block
        r = self.upload('section_pupils', text, commit=True, skip=True)
        self.assertTrue(r.json()['committed'])
        self.assertIn(self.loose, self.sec_b.pupils.all())
        self.assertFalse(User.objects.filter(email='ghost@x.edu').exists())  # never creates accounts
        self.assertEqual(self.upload('section_pupils', 'Email,Add to section\na1@x.edu,\n').status_code, 400)

    # ---------------- bulk: pupils into subjects (whole sections or individuals)
    def test_bulk_add_pupils_to_subjects_by_section_and_by_pupil(self):
        eco = Subject.objects.create(school=self.school, name='Economics')
        phy = Subject.objects.create(school=self.school, name='Physics')
        eco.pupils.add(self.a2)
        self.client.force_login(self.admin)
        text = ('Email,Pupil name,Section,Subjects\n'
                ',(everyone in this section),Pre-U 1 - A,Economics; Physics\n'
                'b1@x.edu,Ben,Pre-U 1 - B,Economics\n'
                'n1@x.edu,Nia,,\n'                          # blank → skipped
                ',,,Economics\n'                             # neither email nor section
                ',,Pre-U 1 - A,Chemistry\n')                 # unknown subject
        r = self.upload('subject_pupils', text)
        res = r.json()
        self.assertEqual(res['summary'], {'create': 0, 'update': 2, 'unchanged': 0, 'error': 2}, res['rows'])
        self.assertIn('Economics (+1)', res['rows'][0]['detail'])   # Abe was already in Economics
        self.assertIn('Physics (+2)', res['rows'][0]['detail'])
        self.upload('subject_pupils', text, commit=True, skip=True)
        self.assertEqual(set(eco.pupils.all()), {self.a1, self.a2, self.b1})
        self.assertEqual(set(phy.pupils.all()), {self.a1, self.a2})
        # Re-running is a no-op.
        again = self.upload('subject_pupils', 'Email,Pupil name,Section,Subjects\n,,Pre-U 1 - A,Economics; Physics\n').json()
        self.assertEqual(again['summary']['update'], 0)

    # ---------------- bulk: Roles imports gain a Subjects column
    def test_teacher_and_pupil_imports_accept_subjects(self):
        Subject.objects.create(school=self.school, name='Economics')
        self.client.force_login(self.admin)
        r = self.upload('teachers', 'First name,Email,Sections,Subjects\nEco,eco@x.edu,Pre-U 1 - A,Economics\n', commit=True)
        self.assertTrue(r.json()['committed'], r.content)
        self.assertEqual(list(self.teacher.subjects_taught.values_list('name', flat=True)), ['Economics'])
        self.assertIn(self.teacher, self.sec_a.teachers.all())
        r = self.upload('pupils', 'First name,Email,Section,Subjects\nZed,zed@x.edu,Pre-U 1 - B,Economics\n', commit=True)
        self.assertTrue(r.json()['committed'], r.content)
        zed = User.objects.get(email='zed@x.edu')
        self.assertIn(zed, self.sec_b.pupils.all())
        self.assertEqual(list(zed.subjects_enrolled.values_list('name', flat=True)), ['Economics'])
        bad = self.upload('pupils', 'First name,Email,Subjects\nYo,yo@x.edu,Chemistry\n').json()
        self.assertEqual(bad['summary']['error'], 1)
        self.assertIn('Subjects page', bad['rows'][0]['message'])

    # ---------------- templates round-trip
    def test_templates_download_and_reupload(self):
        from openpyxl import load_workbook
        eco = Subject.objects.create(school=self.school, name='Economics')
        self.client.force_login(self.admin)
        for kind in ('classes', 'subjects', 'section_pupils', 'subject_pupils', 'pupils', 'teachers'):
            r = self.client.get(f'/s/ssc/api/bulk/{kind}/template/')
            self.assertEqual(r.status_code, 200, kind)
            self.assertIn('Instructions', load_workbook(io.BytesIO(r.content)).sheetnames)
        # Fill in the subject-pupils template: Economics for the whole of section B (its shaded row).
        wb = load_workbook(io.BytesIO(self.client.get('/s/ssc/api/bulk/subject_pupils/template/').content))
        ws = wb['Subject pupils']
        row = next(r for r in range(2, ws.max_row + 1) if ws.cell(r, 4).value == 'Pre-U 1 - B' and not ws.cell(r, 3).value)
        ws.cell(row, 5, 'Economics')
        out = io.BytesIO(); wb.save(out)
        up = SimpleUploadedFile('subject-pupils.xlsx', out.getvalue())
        r = self.client.post('/s/ssc/api/bulk/subject_pupils/import/', {'file': up, 'commit': '1'})
        self.assertTrue(r.json()['committed'], r.content)
        self.assertEqual(list(eco.pupils.all()), [self.b1])
        # Section-pupils template lists every pupil with their current section.
        wb = load_workbook(io.BytesIO(self.client.get('/s/ssc/api/bulk/section_pupils/template/').content))
        listed = [r for r in wb['Year group pupils'].iter_rows(min_row=2, values_only=True) if r[2]]
        rows = {r[2]: r[3] for r in listed}
        # Listed by last name A–Z, with first and last name in their own columns.
        self.assertEqual(wb['Year group pupils']['A1'].value, 'First name')
        self.assertEqual([r[1] for r in listed], sorted((r[1] for r in listed), key=lambda v: (v or '').lower()))
        self.assertEqual(rows['a1@x.edu'], 'Pre-U 1 - A')
        self.assertIsNone(rows['n1@x.edu'])
        self.assertNotIn('f1@x.edu', rows)

    def test_assign_picker_loads_subjects_for_teacher(self):
        eco = Subject.objects.create(school=self.school, name='Economics')
        eco.teachers.add(self.teacher)
        eco.pupils.add(self.a1, self.b1)
        self.client.force_login(self.teacher)
        page = self.client.get('/s/ssc/').content.decode()
        self.assertIn("request('/api/subjects/')", page)
        r = self.client.post('/s/ssc/api/assignments/create/', {
            'title': 'Quiz', 'items_meta': json.dumps([{'type': 'long_answer', 'label': 'Q1', 'prompt': 'Why?'}]),
            'recipients': json.dumps([{'name': 'Ben', 'email': 'b1@x.edu'}]), 'visibility': 'now'})
        self.assertIn(r.status_code, (200, 201), r.content)
        r = self.client.post('/s/ssc/api/assignments/create/', {
            'title': 'Quiz', 'items_meta': json.dumps([{'type': 'long_answer', 'label': 'Q1', 'prompt': 'Why?'}]),
            'recipients': json.dumps([{'name': 'Nia', 'email': 'n1@x.edu'}]), 'visibility': 'now'})
        self.assertEqual(r.status_code, 403)


class CommaSeparatedListsTests(TestCase):
    """Several sections / subjects in one cell are separated by commas (";" still accepted)."""
    def setUp(self):
        self.school = School.objects.create(name='SSC', slug='ssc')
        self.admin = User.objects.create_user('ad', 'ad@x.edu', 'x', school=self.school, role='teacher', is_school_admin=True)
        self.a = SchoolClass.objects.create(school=self.school, name='Pre-U 1 - A')
        self.b = SchoolClass.objects.create(school=self.school, name='Pre-U 1 - B')
        self.odd = SchoolClass.objects.create(school=self.school, name='Clubs, Chess')   # a name with a comma in it
        self.eco = Subject.objects.create(school=self.school, name='Economics')
        self.phy = Subject.objects.create(school=self.school, name='Physics')
        self.p = User.objects.create_user('p', 'p@x.edu', 'x', school=self.school, first_name='Pia')
        self.client.force_login(self.admin)

    def upload(self, kind, text):
        return self.client.post(f'/s/ssc/api/bulk/{kind}/import/', {'file': csv_file(text), 'commit': '1'}).json()

    def test_commas_split_sections_and_subjects(self):
        self.upload('section_pupils', 'Email,Add to section\np@x.edu,"Pre-U 1 - A, Pre-U 1 - B"\n')
        self.assertEqual(set(self.p.classes_enrolled.all()), {self.a, self.b})
        self.upload('subject_pupils', 'Email,Section,Subjects\np@x.edu,,"Economics, Physics"\n')
        self.assertEqual(set(self.p.subjects_enrolled.all()), {self.eco, self.phy})
        r = self.upload('pupils', 'First name,Email,Section,Subjects\nZed,z@x.edu,"Pre-U 1 - A, Pre-U 1 - B",Economics\n')
        self.assertTrue(r['committed'], r)
        self.assertEqual(User.objects.get(email='z@x.edu').classes_enrolled.count(), 2)

    def test_whole_cell_matching_a_name_with_a_comma_is_not_split(self):
        self.upload('section_pupils', 'Email,Add to section\np@x.edu,"Clubs, Chess"\n')
        self.assertEqual(list(self.p.classes_enrolled.all()), [self.odd])

    def test_templates_use_commas(self):
        from openpyxl import load_workbook
        self.a.pupils.add(self.p); self.b.pupils.add(self.p)
        wb = load_workbook(io.BytesIO(self.client.get('/s/ssc/api/bulk/section_pupils/template/').content))
        row = next(r for r in wb['Year group pupils'].iter_rows(min_row=2, values_only=True) if r[2] == 'p@x.edu')
        self.assertEqual(row[3], 'Pre-U 1 - A, Pre-U 1 - B')
        text = ' '.join(str(c) for r in wb['Instructions'].iter_rows(values_only=True) for c in r if c)
        self.assertIn('commas', text)
        self.assertNotIn('";"', text)


class PrefilledPeopleTemplatesTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name='SSC', slug='ssc')
        self.admin = User.objects.create_user('ad', 'ad@x.edu', 'x', school=self.school, role='teacher',
                                              is_school_admin=True, first_name='Mik', last_name='Yang')
        self.t = User.objects.create_user('t', 't@x.edu', 'x', school=self.school, role='teacher', first_name='Ana', last_name='Reyes')
        User.objects.create_superuser('root', 'root@x.edu', 'x', school=self.school)
        self.y7 = SchoolClass.objects.create(school=self.school, name='Year 7')
        self.eco = Subject.objects.create(school=self.school, name='Economics')
        self.eco.teachers.add(self.t)
        self.p1 = User.objects.create_user('p1', 'p1@x.edu', 'x', school=self.school, first_name='Zoe', last_name='Abad')
        self.p2 = User.objects.create_user('p2', 'p2@x.edu', 'x', school=self.school, first_name='Ben', last_name='Cruz')
        self.y7.pupils.add(self.p1); self.eco.pupils.add(self.p1)
        self.client.force_login(self.admin)

    def rows(self, kind, prefill):
        from openpyxl import load_workbook
        r = self.client.get(f'/s/ssc/api/bulk/{kind}/template/' + ('?prefill=1' if prefill else ''))
        self.assertEqual(r.status_code, 200)
        return r, list(load_workbook(io.BytesIO(r.content)).worksheets[0].iter_rows(min_row=2, values_only=True))

    def test_blank_and_prefilled_downloads(self):
        _, blank = self.rows('pupils', False)
        self.assertEqual([r for r in blank if any(r)], [])
        resp, pupils = self.rows('pupils', True)
        self.assertIn('current', resp['Content-Disposition'])
        self.assertEqual(pupils, [('Zoe', 'Abad', 'p1@x.edu', 'Year 7', 'Economics'), ('Ben', 'Cruz', 'p2@x.edu', None, None)])
        _, staff = self.rows('teachers', True)
        self.assertEqual(staff, [('Ana', 'Reyes', 't@x.edu', 'Teacher', 'Economics'), ('Mik', 'Yang', 'ad@x.edu', 'Admin', None)])
        _, blank_staff = self.rows('teachers', False)
        self.assertEqual([r for r in blank_staff if any(r)], [])

    def test_reuploading_prefilled_files_changes_nothing(self):
        for kind in ('pupils', 'teachers'):
            r = self.client.get(f'/s/ssc/api/bulk/{kind}/template/?prefill=1')
            up = SimpleUploadedFile(f'{kind}.xlsx', r.content)
            res = self.client.post(f'/s/ssc/api/bulk/{kind}/import/', {'file': up, 'commit': '0'}).json()
            self.assertEqual(res['summary']['update'] + res['summary']['create'] + res['summary']['error'], 0, (kind, res['rows']))
            self.assertEqual(res['summary']['unchanged'], 2)


class ExamOwnershipTests(TestCase):
    """Plain teachers see and manage only their own exams; admins see every exam."""
    def setUp(self):
        from .models import Assignment, Recipient
        self.school = School.objects.create(name='SSC', slug='ssc')
        self.admin = User.objects.create_user('ad', 'ad@x.edu', 'x', school=self.school, role='teacher', is_school_admin=True)
        self.ana = User.objects.create_user('ana', 'ana@x.edu', 'x', school=self.school, role='teacher', first_name='Ana')
        self.ben = User.objects.create_user('ben', 'ben@x.edu', 'x', school=self.school, role='teacher', first_name='Ben')
        self.p = User.objects.create_user('p', 'p@x.edu', 'x', school=self.school, first_name='Pia')
        self.y7 = SchoolClass.objects.create(school=self.school, name='Year 7')
        self.y7.pupils.add(self.p); self.y7.teachers.add(self.ben)      # Ben is on the year group only
        eco = Subject.objects.create(school=self.school, name='Economics')
        eco.teachers.add(self.ana); eco.pupils.add(self.p)
        self.exam = Assignment.objects.create(school=self.school, title='Ana quiz', teacher_name='Ana', teacher_email='ana@x.edu')
        self.rec = Recipient.objects.create(assignment=self.exam, pupil_name='Pia', pupil_email='p@x.edu')

    def post(self, url, data=None):
        return self.client.post('/s/ssc' + url, json.dumps(data or {}), content_type='application/json')

    def titles(self):
        return [a['title'] for a in self.client.get('/s/ssc/api/assignments/').json()['assignments']]

    def test_list_is_scoped(self):
        self.client.force_login(self.ana); self.assertEqual(self.titles(), ['Ana quiz'])
        self.client.force_login(self.ben); self.assertEqual(self.titles(), [])
        self.client.force_login(self.admin); self.assertEqual(self.titles(), ['Ana quiz'])

    def test_other_teacher_cannot_touch_the_exam(self):
        self.client.force_login(self.ben)
        e, r = self.exam.id, self.rec.id
        self.assertEqual(self.client.get(f'/s/ssc/api/assignments/{e}/status/').status_code, 403)
        self.assertEqual(self.post(f'/api/assignments/{e}/update/', {'title': 'Hacked'}).status_code, 403)
        self.assertEqual(self.post(f'/api/assignments/{e}/state/', {'action': 'close'}).status_code, 403)
        self.assertEqual(self.post(f'/api/assignments/{e}/delete/', {'password': 'x'}).status_code, 403)
        self.assertEqual(self.post(f'/api/assignments/{e}/broadcast-message/', {'text': 'hi'}).status_code, 403)
        self.assertEqual(self.post(f'/api/assignments/{e}/publish-grades/').status_code, 403)
        self.assertEqual(self.client.get(f'/s/ssc/api/recipients/{r}/messages/').status_code, 403)
        self.assertEqual(self.post(f'/api/recipients/{r}/messages/send/', {'text': 'hi'}).status_code, 403)
        self.assertEqual(self.post(f'/api/recipients/{r}/accommodations/', {'extraMinutes': 5}).status_code, 403)
        self.assertEqual(self.post(f'/api/submissions/{r}/grade/', {'itemScores': {}}).status_code, 403)
        self.assertEqual(self.post(f'/api/submissions/{r}/unlock/').status_code, 403)
        self.exam.refresh_from_db(); self.assertEqual(self.exam.title, 'Ana quiz')
        # The owner and an admin can.
        for who in (self.ana, self.admin):
            self.client.force_login(who)
            self.assertEqual(self.client.get(f'/s/ssc/api/assignments/{e}/status/').status_code, 200)

    def test_teacher_assigns_by_subject_not_year_group(self):
        self.client.force_login(self.ben)       # year group only → no pupils to assign to
        self.assertEqual(self.client.get('/s/ssc/api/pupils/').json()['pupils'], [])
        self.client.force_login(self.ana)       # subject teacher
        self.assertEqual([p['email'] for p in self.client.get('/s/ssc/api/pupils/').json()['pupils']], ['p@x.edu'])
        page = self.client.get('/s/ssc/').content.decode()
        self.assertIn('Tick a subject to select everyone in it', page)
