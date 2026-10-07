import json, os, shutil, tempfile
from django.core.files.base import ContentFile
from django.core.management import call_command
from django.test import TestCase, override_settings
from .models import Assignment, AssignmentItem, Recipient, School, SchoolClass, StoredFile, Submission, User

MEDIA = tempfile.mkdtemp()


@override_settings(MEDIA_ROOT=MEDIA)
class FilesSurviveAcrossServersTests(TestCase):
    @classmethod
    def tearDownClass(cls):
        super().tearDownClass(); shutil.rmtree(MEDIA, ignore_errors=True)

    def setUp(self):
        self.school = School.objects.create(name='SSC', slug='ssc')
        self.t = User.objects.create_user('t', 't@x.edu', 'x', school=self.school, role='teacher', is_school_admin=True)
        self.a = Assignment.objects.create(school=self.school, title='E', teacher_name='T', teacher_email='t@x.edu')
        self.item = AssignmentItem.objects.create(assignment=self.a, item_type='pdf', label='Text 1')
        self.item.pdf.save('doc.pdf', ContentFile(b'%PDF-1.4 hello'))

    def test_upload_is_copied_to_db_and_served_when_disk_copy_is_gone(self):
        self.assertTrue(StoredFile.objects.filter(name=self.item.pdf.name).exists())
        os.remove(os.path.join(MEDIA, self.item.pdf.name))  # e.g. Railway redeploy / other machine
        self.client.force_login(self.t)
        r = self.client.get(f'/s/ssc/api/items/{self.item.id}/pdf/')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(b''.join(r.streaming_content), b'%PDF-1.4 hello')
        self.assertTrue(os.path.exists(os.path.join(MEDIA, self.item.pdf.name)))  # re-cached

    def test_missing_everywhere_gives_clear_404(self):
        StoredFile.objects.all().delete(); os.remove(os.path.join(MEDIA, self.item.pdf.name))
        self.client.force_login(self.t)
        r = self.client.get(f'/s/ssc/api/items/{self.item.id}/pdf/')
        self.assertEqual(r.status_code, 404); self.assertEqual(r.json()['code'], 'file_missing')

    def test_sync_command_copies_disk_only_files(self):
        StoredFile.objects.all().delete()
        call_command('sync_media_to_db', verbosity=0, stdout=open(os.devnull, 'w'))
        self.assertTrue(StoredFile.objects.filter(name=self.item.pdf.name).exists())


class EditDeletePeopleTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name='SSC', slug='ssc')
        self.other = School.objects.create(name='B', slug='b')
        self.admin = User.objects.create_user('ad', 'ad@x.edu', 'x', school=self.school, role='teacher', is_school_admin=True)
        self.teacher = User.objects.create_user('t', 't@x.edu', 'x', school=self.school, role='teacher', first_name='Tess')
        self.p1 = User.objects.create_user('p1', 'p1@x.edu', 'x', school=self.school, first_name='Ana')
        self.p2 = User.objects.create_user('p2', 'p2@x.edu', 'x', school=self.school, first_name='Ben')
        self.far = User.objects.create_user('p3', 'p3@x.edu', 'x', school=self.other)
        self.exam = Assignment.objects.create(school=self.school, title='E', teacher_name='Tess', teacher_email='t@x.edu')
        self.rec = Recipient.objects.create(assignment=self.exam, pupil_name='Ana', pupil_email='p1@x.edu')
        Submission.objects.create(recipient=self.rec, status='submitted')

    def post(self, url, data):
        return self.client.post('/s/ssc' + url, json.dumps(data), content_type='application/json')

    def test_edit_pupil_updates_exam_links(self):
        self.client.force_login(self.admin)
        r = self.post(f'/api/accounts/{self.p1.id}/update/', {'firstName': 'Anna', 'lastName': 'Cruz', 'email': 'anna@x.edu'})
        self.assertEqual(r.status_code, 200, r.content)
        self.rec.refresh_from_db()
        self.assertEqual((self.rec.pupil_email, self.rec.pupil_name), ('anna@x.edu', 'Anna Cruz'))

    def test_edit_teacher_updates_exam_owner_and_rejects_duplicate_email(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.post(f'/api/accounts/{self.teacher.id}/update/', {'firstName': 'T', 'email': 'p2@x.edu'}).status_code, 400)
        self.post(f'/api/accounts/{self.teacher.id}/update/', {'firstName': 'Tess', 'lastName': 'Reyes', 'email': 'tess@x.edu'})
        self.exam.refresh_from_db(); self.assertEqual(self.exam.teacher_email, 'tess@x.edu')

    def test_delete_pupils_keeps_their_work(self):
        self.client.force_login(self.admin)
        r = self.post('/api/accounts/pupils/delete/', {'userIds': [self.p1.id, self.p2.id]})
        self.assertEqual(r.json()['deleted'], 2)
        self.assertFalse(User.objects.filter(id__in=[self.p1.id, self.p2.id]).exists())
        self.assertTrue(Submission.objects.filter(recipient=self.rec).exists())

    def test_delete_guards(self):
        self.client.force_login(self.admin)
        self.assertEqual(self.post('/api/accounts/pupils/delete/', {'userIds': [self.teacher.id]}).status_code, 400)
        self.assertEqual(self.post('/api/accounts/pupils/delete/', {'userIds': [self.far.id]}).status_code, 400)
        self.client.force_login(self.teacher)
        self.assertEqual(self.post('/api/accounts/pupils/delete/', {'userIds': [self.p1.id]}).status_code, 403)
        self.assertEqual(self.post(f'/api/accounts/{self.p1.id}/update/', {'firstName': 'X', 'email': 'x@x.edu'}).status_code, 403)

    def test_delete_staff(self):
        other_admin = User.objects.create_user('ad2', 'ad2@x.edu', 'x', school=self.school, role='teacher', is_school_admin=True)
        boss = User.objects.create_user('root', 'root@x.edu', 'x', school=self.school, role='teacher', is_superuser=True)
        far_teacher = User.objects.create_user('ft', 'ft@x.edu', 'x', school=self.other, role='teacher')
        self.client.force_login(self.admin)
        url = '/api/accounts/staff/delete/'
        self.assertEqual(self.post(url, {'userIds': [self.admin.id]}).status_code, 400)       # not yourself
        self.assertEqual(self.post(url, {'userIds': [boss.id]}).status_code, 400)             # not a main admin
        self.assertEqual(self.post(url, {'userIds': [self.p1.id]}).status_code, 400)          # not a pupil
        self.assertEqual(self.post(url, {'userIds': [far_teacher.id]}).status_code, 400)      # not another school
        r = self.post(url, {'userIds': [self.teacher.id, other_admin.id]})
        self.assertEqual(r.json()['deleted'], 2)
        self.assertFalse(User.objects.filter(id__in=[self.teacher.id, other_admin.id]).exists())
        self.assertTrue(type(self.exam).objects.filter(id=self.exam.id).exists())            # their exam stays
        self.client.force_login(boss)
        self.assertEqual(self.client.post('/s/ssc' + url, '{"userIds": [%d]}' % self.admin.id,
                                          content_type='application/json').json()['deleted'], 1)
        plain = User.objects.create_user('pl', 'pl@x.edu', 'x', school=self.school, role='teacher')
        self.client.force_login(plain)
        self.assertEqual(self.post(url, {'userIds': [far_teacher.id]}).status_code, 403)
        self.client.force_login(boss)
        roles = self.client.get('/s/ssc/roles/').content.decode()
        self.assertIn('/api/accounts/staff/delete/', roles)

    def test_pupil_page_starts_in_pupil_layout(self):
        self.client.force_login(self.p1)
        html = self.client.get('/s/ssc/').content.decode()
        self.assertIn('<div id="teacherView" class="hidden">', html)
        self.assertIn('id="pupilLoading"', html)
        self.assertIn('<div id="studentLogin" class="hidden">', html)
        self.assertIn('class="role-toggle hidden"', html)


class NameEditAndSuperadminTests(TestCase):
    def setUp(self):
        self.school = School.objects.create(name='SSC', slug='ssc')
        self.admin = User.objects.create_user('ad', 'ad@x.edu', 'x', school=self.school, role='teacher', is_school_admin=True)
        self.p = User.objects.create_user('p1', 'shared@x.edu', 'x', school=self.school, first_name='Ana')
        # An old test login elsewhere that happens to use the same email.
        User.objects.create_user('old', 'Shared@x.edu', 'x', role='teacher')

    def test_editing_only_the_name_works_even_if_email_is_shared(self):
        self.client.force_login(self.admin)
        r = self.client.post(f'/s/ssc/api/accounts/{self.p.id}/update/',
                             json.dumps({'firstName': 'Anna', 'lastName': 'Cruz', 'email': 'shared@x.edu'}), content_type='application/json')
        self.assertEqual(r.status_code, 200, r.content)
        self.p.refresh_from_db(); self.assertEqual(self.p.get_full_name(), 'Anna Cruz')
        r = self.client.post(f'/s/ssc/api/accounts/{self.admin.id}/update/',
                             json.dumps({'firstName': 'A', 'email': 'shared@x.edu'}), content_type='application/json')
        self.assertEqual(r.status_code, 400)   # changing TO a taken email is still refused

    def test_superadmin_edits_school_admin_name_and_email(self):
        boss = User.objects.create_superuser('boss', 'boss@x.edu', 'x')
        exam = Assignment.objects.create(school=self.school, title='E', teacher_name='ad', teacher_email='ad@x.edu')
        self.client.force_login(boss)
        url = f'/superadmin/api/schools/{self.school.id}/admins/{self.admin.id}/update/'
        r = self.client.post(url, json.dumps({'first_name': 'Mik', 'last_name': 'Yang', 'email': 'mik@x.edu'}), content_type='application/json')
        self.assertEqual(r.status_code, 200, r.content)
        self.admin.refresh_from_db(); exam.refresh_from_db()
        self.assertEqual((self.admin.first_name, self.admin.last_name, self.admin.email), ('Mik', 'Yang', 'mik@x.edu'))
        self.assertEqual((exam.teacher_email, exam.teacher_name), ('mik@x.edu', 'Mik Yang'))
        self.assertEqual(self.client.post(url, json.dumps({'first_name': 'M', 'email': 'boss@x.edu'}),
                                          content_type='application/json').status_code, 400)
        listed = self.client.get(f'/superadmin/api/schools/{self.school.id}/admins/').json()
        self.assertEqual(listed[0]['last_name'], 'Yang')
        self.client.force_login(self.admin)
        self.assertNotEqual(self.client.post(url, json.dumps({'first_name': 'X', 'email': 'x@x.edu'}),
                                             content_type='application/json').status_code, 200)

    def test_teachers_template_has_subjects_not_year_groups(self):
        import io
        from openpyxl import load_workbook
        self.client.force_login(self.admin)
        wb = load_workbook(io.BytesIO(self.client.get('/s/ssc/api/bulk/teachers/template/').content))
        headers = [c.value for c in wb['Teachers'][1]]
        self.assertEqual(headers, ['First name *', 'Last name', 'Email *', 'Role', 'Subjects'])
        self.assertIn('Subject list', wb.sheetnames)
        self.assertNotIn('Year group list', wb.sheetnames)
